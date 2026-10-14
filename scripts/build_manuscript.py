#!/usr/bin/env python3
"""Build the study manuscript from verified complete study outputs.

Run with the bundled document runtime, which supplies python-docx.
"""
import json
from pathlib import Path

from docx import Document
from docx.shared import Inches, Pt, RGBColor
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.enum.table import WD_TABLE_ALIGNMENT, WD_CELL_VERTICAL_ALIGNMENT
from docx.enum.style import WD_STYLE_TYPE
from docx.oxml import OxmlElement
from docx.oxml.ns import qn

ROOT = Path(__file__).resolve().parents[1]
DATA = json.loads((ROOT/'results/study/analysis/manuscript_data.json').read_text())
PATIENT = json.loads((ROOT/'results/study/analysis/patient_uncertainty.json').read_text())
S = DATA['summary']
ADDITIONAL = None
combined = ROOT/'results/combined/summary.json'
if combined.exists():
    value = json.loads(combined.read_text())
    if value['status'] == 'complete' and not value['smoke']:
        ADDITIONAL = value
        DATA = json.loads((ROOT/'results/combined/analysis/manuscript_data.json').read_text())
        S = DATA['summary']
        PATIENT = json.loads((ROOT/'results/combined/analysis/patient_uncertainty.json').read_text())
assert S['status'] == 'complete' and not S['smoke']
assert all(len(v['completed']) == 10 and not v['missing'] for v in S['coverage'].values())
METHODS = ['supervised','simclr_generic','simclr_domain','byol_generic','byol_domain','clocs_cmsc']
NAMES = {'supervised':'Supervised focal + OS', 'simclr_generic':'SimCLR generic',
         'simclr_domain':'SimCLR domain', 'byol_generic':'BYOL generic',
         'byol_domain':'BYOL domain', 'clocs_cmsc':'CLOCS CMSC adaptation'}
for n in METHODS[1:]:
    NAMES[n+'_finetuned']=NAMES[n]+' fine-tuned'
NAMES.update(clocs_reference_supervised='Reference CLOCS supervised',
             clocs_reference_cmsc='Reference CLOCS frozen',
             clocs_reference_finetuned='Reference CLOCS fine-tuned')
DISPLAY_METHODS=METHODS+[n+'_finetuned' for n in METHODS[1:]]+['clocs_reference_supervised','clocs_reference_cmsc','clocs_reference_finetuned']


def interval(s):
    return f"[{s['ci95_mean'][0]:.4f}, {s['ci95_mean'][1]:.4f}]"


def estimate(s):
    return f"{s['mean']:.4f} ± {s['sample_sd']:.4f}"


doc = Document('/Users/tumokgabeng/Documents/UNISA/Deep Learning Indaba/Deep Learning Indaba.docx')
# Construct a consistent scientific article body and section layout.
body = doc._element.body
for child in list(body):
    body.remove(child)
sect = OxmlElement('w:sectPr')
body.append(sect)
section = doc.sections[0]
section.page_width, section.page_height = Inches(8.27), Inches(11.69)
section.top_margin = section.bottom_margin = Inches(.7)
section.left_margin = section.right_margin = Inches(.75)
section.header_distance = section.footer_distance = Inches(.3)
for name in ['Normal','Title','Heading 1','Heading 2','Caption']:
    try:
        st = doc.styles[name]
    except KeyError:
        st = doc.styles.add_style(name, WD_STYLE_TYPE.PARAGRAPH)
    st.font.name = 'Times New Roman'
    st.font.color.rgb = RGBColor(0,0,0)
    st.paragraph_format.space_before = Pt(0)
    st.paragraph_format.space_after = Pt(6)
    st.paragraph_format.line_spacing = 1.08
doc.styles['Normal'].font.size = Pt(10.5)
doc.styles['Title'].font.size = Pt(16)
doc.styles['Title'].font.bold = True
doc.styles['Heading 1'].font.size = Pt(12)
doc.styles['Heading 2'].font.size = Pt(11)
doc.styles['Caption'].font.size = Pt(9)
for name in ['Heading 1','Heading 2']:
    doc.styles[name].font.bold = True
    doc.styles[name].paragraph_format.space_before = Pt(9)
    doc.styles[name].paragraph_format.keep_with_next = True


def paragraph(text, style=None):
    p = doc.add_paragraph(text, style)
    p.paragraph_format.widow_control = True
    return p


def heading(text, level=1):
    doc.add_heading(text, level)


def figure(name, caption):
    paragraph(caption, 'Caption').paragraph_format.keep_with_next = True
    doc.add_picture(str(ROOT/'results/combined/analysis'/name), width=Inches(6.7))


def table(caption, headers, rows, widths=None):
    p = paragraph(caption, 'Caption')
    p.paragraph_format.keep_with_next = True
    t = doc.add_table(rows=1, cols=len(headers))
    t.alignment = WD_TABLE_ALIGNMENT.CENTER
    t.autofit = False
    if widths is None:
        widths = [6.77/len(headers)]*len(headers)
    for col, width in zip(t.columns,widths):
        col.width = Inches(width)
    for cell, text in zip(t.rows[0].cells,headers):
        cell.text = text
    for row in rows:
        for cell, text in zip(t.add_row().cells,row):
            cell.text = str(text)
    for ri,row in enumerate(t.rows):
        props = row._tr.get_or_add_trPr()
        props.append(OxmlElement('w:cantSplit'))
        if ri == 0:
            props.append(OxmlElement('w:tblHeader'))
        for ci,cell in enumerate(row.cells):
            cell.width = Inches(widths[ci])
            cell.vertical_alignment = WD_CELL_VERTICAL_ALIGNMENT.CENTER
            tcpr = cell._tc.get_or_add_tcPr()
            borders = OxmlElement('w:tcBorders')
            for edge in ['top','left','bottom','right']:
                e = OxmlElement('w:'+edge)
                for key,value in [('val','single'),('sz','4'),('color','D9D9D9')]:
                    e.set(qn('w:'+key),value)
                borders.append(e)
            tcpr.append(borders)
            margin = OxmlElement('w:tcMar')
            for edge in ['top','left','bottom','right']:
                e=OxmlElement('w:'+edge); e.set(qn('w:w'),'65');e.set(qn('w:type'),'dxa');margin.append(e)
            tcpr.append(margin)
            if ri == 0:
                shade=OxmlElement('w:shd');shade.set(qn('w:fill'),'E7E7E7');tcpr.append(shade)
            for p in cell.paragraphs:
                p.alignment = WD_ALIGN_PARAGRAPH.LEFT if ci == 0 else WD_ALIGN_PARAGRAPH.CENTER
                p.paragraph_format.space_after = Pt(0)
                p.paragraph_format.line_spacing = 1.0
                p.paragraph_format.keep_with_next = ri < len(t.rows)-1
                for run in p.runs:
                    run.font.size = Pt(9)
                    run.font.bold = ri == 0
                    run.font.color.rgb=RGBColor(0,0,0)
    paragraph('')
    return t


title=paragraph('Evaluating domain adaptive augmentations for self supervised ECG classification','Title')
title.alignment=WD_ALIGN_PARAGRAPH.CENTER
heading('Abstract')
paragraph('Self-supervised ECG learning may reduce annotation requirements, but augmentation strength and downstream adaptation require separate controls. We evaluate 570 runs across 57 conditions and ten training seeds: 520 PTB-XL runs with patient-separated folds and 50 MIT-BIH binary annotation-window transfer runs. Controls include generic/domain policies, seven omissions per framework, gate/RMS matched generic policies, frozen and fully fine-tuned encoders, five labelled fractions and architecture-matched CLOCS comparisons. At 10% labels, full domain fine-tuning reaches macro-F1 0.6061 for SimCLR and 0.6067 for BYOL versus supervised 0.5802; paired gains survive the dedicated fine-tuning family (Holm p = 0.0195), but not the broader label-efficiency family. Domain AUROC advantages over supervised learning are not established. Matching generic gate events and perturbation RMS weakens attribution to domain-specific transformations: frozen SimCLR domain AUROC is lower than matched generic AUROC (difference −0.0120; adjusted p = 0.0156), whereas frozen BYOL retains domain F1 and AUROC advantages. Neither framework establishes a fine-tuned advantage over matched generic controls. No label-efficiency or MIT-BIH contrast survives panel adjustment. Omission effects and frozen BYOL’s AUROC deficit support metric- and framework-dependent conclusions. The findings do not establish universal SSL superiority or clinical fidelity.')
paragraph('Keywords: self-supervised learning; ECG; augmentation ablation; label scarcity; patient-separated evaluation')

heading('1 Introduction')
paragraph('The 12-lead electrocardiogram provides an accessible signal for cardiovascular diagnosis, but reliable labels require clinical interpretation. Self-supervised learning (SSL) uses unlabelled training signals to learn an encoder before downstream classification. SimCLR [1] learns through contrastive views, while BYOL [2] matches online predictions to an evolving target representation. Both objectives depend on the views supplied during pretraining.')
paragraph('Augmentations inspired by ECG acquisition and temporal variation may improve transfer, but a comparison against supervised learning alone cannot attribute any difference specifically to the augmentation policy. We therefore separate two questions: whether a domain policy improves SSL over a generic time-series policy, and whether those representations outperform a supervised model with the same labelled training budget. We also test the marginal contribution of each transformation through leave-one-out ablations.')
paragraph('This study evaluates 570 complete runs across 57 conditions and ten training seeds. PTB-XL accounts for 520 runs, including generic/domain controls, seven omissions per framework, full encoder adaptation, CLOCS controls, augmentation matching and five labelled-data fractions. Fifty MIT-BIH runs assess transfer from four PTB-XL encoders against random initialization. We report uncertainty and negative findings without assuming that physiological inspiration guarantees realism or clinical benefit.')

heading('2 Related work')
paragraph('SimCLR [1] introduced a contrastive framework in which view composition and projection heads influence representation quality. BYOL [2] learns from an online network and a moving target without explicit negative pairs. CLOCS [3] developed contrastive cardiac-signal objectives across space, time and patients. Its CMSC temporal-view objective is a relevant signal-only comparator for this study; we adapt it to the common encoder and training budget rather than claiming an exact architectural reproduction.')
paragraph('Other ECG approaches use different inputs or supervision. Fine-grained ECG-report alignment [4] uses paired reports, which are not inputs to our waveform-only protocol. Channel self-attention for varied-lead diagnosis [5] addresses ECG classification with a different architecture and setting. We retain these distinctions instead of treating all ECG-related methods as interchangeable SSL baselines. PTB-XL and its recommended folds provide the data foundation [6, 7]. Focal loss [8] and record oversampling address imbalance in the downstream task.')

heading('3 Methods')
heading('3 1 Dataset and labelled budget',2)
paragraph('We use PTB-XL v1.0.3, containing 21,799 ten-second, 12-lead records from 18,869 patients. The recommended folds 1–8 form training, fold 9 validation and fold 10 test. Patient identifiers are disjoint across the three partitions. SSL uses all 17,418 training records without diagnostic labels, including records lacking a diagnostic superclass. It never uses validation or test waveforms. Diagnostic eligibility yields 17,084 training, 2,146 validation and 2,158 test records.')
paragraph('The five multilabel targets are NORM, CD, HYP, MI and STTC. For the 10% comparison, iterative multilabel stratification for each seed followed by deterministic budget correction selects exactly 1,708 eligible training records, approximately 10% of that pool. The same labelled subset is shared across conditions within each fraction and seed. Random seeds are 42, 52, 62, 72, 82, 92, 102, 112, 122 and 132. Training label composition and model initialization vary with seed; patient folds remain fixed.')
paragraph('We use the provided 100 Hz waveforms, with 1,000 samples per lead. Each lead is normalized within each record to zero mean and unit population standard deviation. All 21,799 waveforms were checked for expected shape, sampling rate and finite values before caching. Class-positive supports are not unique patient or record counts because labels can co-occur.')

heading('3 2 Augmentation policies',2)
paragraph('Both generic and domain policies apply Gaussian jitter with standard deviation 0.03 and probability 0.9, plus amplitude scaling uniformly between 0.9 and 1.1 with probability 0.8. The generic policy additionally zero-masks a contiguous 5–20% temporal segment with probability 0.3. It is a prespecified generic time-series control, not a benchmark of every possible generic policy.')
paragraph('The domain policy adds the seven independently gated transformations in Table 1 instead of generic zero-masking. There is no hidden outer probability gate. Mixing uses batch partners without labels; partners and mixup coefficients are shared between the two views to preserve their source-pair identity. The view-specific gates and other perturbations remain stochastic. The implemented policy is versioned as ecg_v2_100hz.')
table('Table 1 Domain policy transformations and independent probabilities',
      ['Transformation','Parameters used at 100 Hz','Probability'],[
      ['Frequency warping','Monotone piecewise-linear time warp; interior anchors perturbed by up to 5%','0.50'],
      ['Mixup','Beta(1,1) coefficient; anchor weight at least 0.5; label-free batch partner','0.40'],
      ['Bandpass variation','Third-order Butterworth, zero-phase; low 0.5–1.5 Hz, high 30–40 Hz','0.30'],
      ['Segment CutMix','Partner segment covering 10–30%; tapered boundaries','0.25'],
      ['Motion artefact','Sinusoid at 0.5–2 Hz; amplitude 0.01–0.05 in normalized units','0.20'],
      ['Channel noise','Independent per-lead noise; SD 0.005–0.02 times lead SD','0.60'],
      ['Temporal dropout','Replace 5–20% segment by interpolation between its endpoints','0.30']], [1.35,4.62,.8])
paragraph('These transformations are motivated by signal variation, filtering and acquisition artefacts, but their clinical fidelity is not established experimentally. Mixup is not constrained to different diagnostic classes and is not asserted to simulate a real comorbid ECG. The bandpass range respects the 50 Hz Nyquist frequency; no device standard is claimed to mandate this range. Motion amplitudes refer to normalized signal units, not millivolts.')

heading('3 3 Encoder and SSL objectives',2)
paragraph('The shared-backbone conditions use ECGEncoder1DCNN with six convolutional blocks, each comprising convolution, batch normalization and ReLU. Global average pooling produces a 256-dimensional representation. Table 2 gives this architecture; SSL projection heads are discarded before downstream evaluation. Three CLOCS reference conditions instead use the native 500 Hz single-lead architecture described in Section 4 10, with their own matched supervised control.')
table('Table 2 Shared encoder architecture for full ten-second inputs',
      ['Block','Channels','Kernel','Stride','Output shape'],[
      ['Input','12','—','—','12 × 1,000'],['1','64','11','2','64 × 500'],
      ['2','64','7','1','64 × 500'],['3','128','7','2','128 × 250'],
      ['4','128','5','1','128 × 250'],['5','256','5','2','256 × 125'],
      ['6','256','3','1','256 × 125'],['Global average pooling','256','—','—','256']], [2.0,1.0,.8,.8,2.17])
paragraph('SimCLR uses a 256 → 2,048 → 128 projection head and NT-Xent at temperature 0.07. The denominator includes all 2B views except the anchor itself, and the opposite view of the same source is the positive. BYOL uses a 256 → 2,048 → 256 projector and a corresponding predictor, with batch normalization in the hidden layers. Its symmetric loss matches each online prediction to the detached target projection of the opposite view. Target encoder and projector parameters start as copies of their online counterparts and use a fixed EMA coefficient of 0.999; target batch-normalization statistics update through target forwards.')
paragraph('The CLOCS CMSC adaptation uses adjacent, nonoverlapping five-second windows from each ten-second ECG and a shared 12-lead encoder and projection head. Its patient-aware cross-view contrastive objective follows the author implementation [3] at temperature 0.1. Patient identity, not disease labels, defines additional same-patient positives. This comparator uses different temporal views and shorter inputs than the generic and domain policies, so it is an adapted prior-method baseline rather than an isolated augmentation comparison.')

heading('3 4 Training and evaluation',2)
paragraph('All SSL pretraining uses 20 epochs, batch size 128, Adam at an initial learning rate of 0.001 and cosine annealing. PTB-XL downstream training uses 20 epochs, batch size 64 and Adam at 0.001. Supervised controls train the encoder and classification head from scratch. Linear probes freeze encoder parameters and batch-normalization state and train only the linear head. Full fine-tuning trains the encoder and head from the corresponding seed-specific pretrained encoder, with identical downstream settings and labelled subsets. Reused encoders are verified by checksum and parent fingerprint; pretraining is not counted twice.')
paragraph('Every PTB-XL downstream condition uses focal loss with gamma 2, positive alpha 0.25 and negative alpha 0.75, plus weighted record sampling with replacement for one dataset-length epoch. Sampling weights average inverse positive-class frequencies over a record’s positive labels; this does not enforce equal class counts in every minibatch. The focal-loss and oversampling configuration is prespecified; this experiment does not establish that it is optimal among all loss and balancing strategies.')
paragraph('The best checkpoint is selected only by validation macro-F1 at threshold 0.5. Its test predictions are evaluated once after selection. Primary performance is test macro-F1 at 0.5; secondary performance is macro-AUROC. Macro PR-AUC, micro sensitivity/specificity, Brier score and class-averaged ECE provide diagnostics. Secondary class thresholds are selected on validation predictions only and transferred unchanged to test; they do not change AUROC, PR-AUC or probability calibration scores.')
table('Table 3 Experimental matrix at 10% labelled PTB-XL data', ['Conditions','Count','Seeds per condition','Pretraining epochs'],[
      ['Supervised focal + oversampling','1','10','None'],
      ['SimCLR generic and domain','2','10','20'],
      ['BYOL generic and domain','2','10','20'],
      ['CLOCS CMSC adaptation','1','10','20'],
      ['SimCLR domain minus one of seven transformations','7','10','20'],
      ['BYOL domain minus one of seven transformations','7','10','20'],
      ['Full fine-tuning of five shared-backbone SSL conditions','5','10','20 reused'],
      ['CLOCS reference supervised, frozen and fine-tuned','3','10','None / 20 / reused']], [3.57,.8,1.3,1.1])
paragraph('Each ablation disables exactly one domain transformation while retaining all other parameters and common jitter/scaling. Every seed independently trains its SSL encoder; fine-tuned and frozen counterparts within a seed reuse that encoder. Labelled records, downstream sampling seeds and classification-head initialization are paired across conditions. The linear-probe suite rotates condition order between seeds; the omission, fine-tuning and reference suite uses a fixed order to satisfy checkpoint dependencies. All runs execute serially on Apple M5 Pro through PyTorch MPS with four CPU threads. Equal epochs do not equalize input durations, update counts across different datasets, FLOPs or elapsed time.')

heading('3 5 Augmentation matching label efficiency and transfer',2)
paragraph('The matched generic policy uses the same per-record gate events as the domain policy, including common jitter/scaling and seven generic slots: noise, scaling, circular shift, zero-masking, nine-sample moving average, linear trend and five-segment permutation. For each training waveform and view, a domain reference determines total waveform RMS displacement; the generic displacement is rescaled to that value. This controls gate frequency and total RMS distance, not morphology, spectrum or clinical fidelity. Domain reference generation adds computation. Frozen and fully fine-tuned SimCLR/BYOL controls contribute 40 runs.')
paragraph('Label efficiency uses exactly 171, 854, 1,708, 3,417 and 8,542 labelled PTB-XL records at 1%, 5%, 10%, 20% and 50%, respectively. Each fraction and seed shares the same sampled records across supervised and fully fine-tuned generic/domain SimCLR/BYOL. All validation and test subjects remain fixed. Four fractions contribute 200 runs; the existing 10% results provide the fifth stratum. The seed-specific pretrained encoders are checksum verified and reused without additional pretraining.')
paragraph('MIT-BIH v1.0.0 [9–11] contains 48 two-channel records from 47 subjects. We validate every record and split by subject with seed 42, assigning records 201 and 202 together: 29 training subjects, nine validation subjects and nine test subjects yield 5,040, 1,262 and 1,602 eligible nonoverlapping ten-second windows. Windows are positive if any annotation is A, a, J, S, V, E or F; N, L, R, e and j are negative annotations. Windows containing paced or unknown symbols /, f or Q, or no accepted annotation, are excluded. The first measured channel is resampled from 360 to 100 Hz, normalized within the window and repeated across twelve channels for the pretrained interface. This is binary annotation-window classification, not an AAMI beat-classification benchmark.')
paragraph('All five MIT-BIH conditions fully train the same encoder and binary head for ten epochs with binary cross entropy, Adam 0.001 and batch size 32. Four conditions initialize from generic/domain PTB-XL SimCLR/BYOL encoders; the supervised control starts randomly. Validation F1 at 0.5 selects checkpoints. Ten training seeds share one fixed subject partition, previously evaluated in a single-seed run; this is corroborative transfer rather than untouched external confirmation, and does not represent ten independent patient partitions.')

heading('3 6 Uncertainty and reproducibility',2)
paragraph('We report means, sample standard deviations with denominator n−1, and percentile 95% bootstrap confidence intervals for the seed mean using 10,000 resamples. Differences use paired seeds, exact two-sided sign-flip tests over 1,024 sign assignments, paired standardized effect size d_z, and Holm adjustment separately for six headline comparisons, seven SimCLR omissions, seven BYOL omissions, ten fine-tuning comparisons and three reference-architecture comparisons, for each metric. Fixed-threshold F1 and AUROC have separate prespecified Holm families. Fine-tuned generic/domain contrasts fall outside these original inferential families; separate post hoc inference is reported in Section 4 11. Bootstrap intervals are unadjusted descriptive intervals; multiplicity-adjusted decisions use the Holm p-values. Seed intervals capture training variability on fixed test patients. Paired patient-cluster and joint patient-and-seed intervals additionally quantify held-out patient sampling, as described in Section 4 6.')
paragraph('Matching, label efficiency and transfer contrasts were declared before their runs after the 10% outcomes had been observed. Their Holm families contain eight, 30 and six contrasts, respectively, separately for F1 and AUROC. Original families remain unchanged. The five-fraction family includes the previously observed 10% stratum; its larger multiplicity burden can change significance relative to the dedicated 10% family. Each panel uses 2,000 paired patient-cluster and joint patient/seed bootstrap draws. All MIT-BIH AUROC draws were valid. These intervals are pointwise and conditional on the training and validation subjects, not adjusted tests.')
paragraph('Each run stores data and source fingerprints, augmentation specifications, labelled splits, checkpoints, epoch histories and validation/test predictions. Resume checkpoints preserve optimizer and random-number-generator state. Tests cover contrastive positive-pair indexing, BYOL targets, augmentation omissions, patient separation, exact label budgets, metric pairing and checkpoint resumption. Pipeline validation records 29 passed tests, 15 additional-condition MPS smoke tests, 14 real-data entry-point checks and seven finite-gradient/checkpoint checks. All prediction metrics are recomputed from saved arrays before reporting. MPS random generators are seeded, but bitwise GPU reproducibility is not guaranteed. The experimental protocol fixes the patient split, normalization, loss definitions and encoder freezing before evaluation.')

heading('4 Results')
heading('4 1 Supervised baseline',2)
paragraph('All 570 declared runs completed, with ten seeds for every condition. The 280-run experiment at 10% labelled PTB-XL data includes the supervised, omission, encoder-adaptation and CLOCS comparisons. The shared-backbone supervised baseline averages 0.5802 macro-F1 and 0.8534 macro-AUROC. Table 4 reports the same uncertainty conventions used for SSL. A seed maximum is not an upper bound on attainable supervised performance, and comparison against two baseline standard deviations is not used as a significance test.')
table('Table 4 Supervised baseline on held-out test records across ten seeds',
      ['Metric','Mean ± sample SD','95% CI of seed mean','Min to max'],[
      [label,estimate(S['summaries']['supervised'][key]),interval(S['summaries']['supervised'][key]),f"{S['summaries']['supervised'][key]['min']:.4f} to {S['summaries']['supervised'][key]['max']:.4f}"]
      for key,label in [('f1_macro','Macro-F1'),('auroc_macro','Macro-AUROC'),('sensitivity_micro','Micro sensitivity'),('specificity_micro','Micro specificity')]], [1.5,1.7,1.9,1.67])

heading('4 2 Generic and domain augmentation controls',2)
table('Table 5 Main test results with symmetric ten-seed reporting',
      ['Method','Macro-F1 mean ± SD','95% CI','AUROC mean ± SD','95% CI'],[
      [NAMES[n],estimate(S['summaries'][n]['f1_macro']),interval(S['summaries'][n]['f1_macro']),estimate(S['summaries'][n]['auroc_macro']),interval(S['summaries'][n]['auroc_macro'])] for n in METHODS], [1.57,1.3,1.3,1.3,1.3])
paragraph('Domain minus generic SimCLR has a paired mean F1 difference of +0.1088 (95% CI [0.0865, 0.1341]) and AUROC difference of +0.0177 ([0.0113, 0.0239]). Domain minus generic BYOL improves F1 by +0.1893 ([0.1725, 0.2083]) and AUROC by +0.0203 ([0.0175, 0.0226]). Holm-adjusted p = 0.0117 for each comparison and metric. This directly supports a benefit from the tested domain policy over its prespecified generic control within each SSL framework.')
paragraph('Both full domain SSL models remain below supervised learning. SimCLR domain minus supervised has F1 difference −0.3197 and AUROC difference −0.0389. BYOL domain minus supervised has F1 difference −0.2684 and AUROC difference −0.0130. All four adjusted p-values are 0.0117. BYOL domain exceeds SimCLR domain by 0.0513 F1 and 0.0259 AUROC. Supervised learning is strongest under this fixed-threshold linear-probe protocol.')
paragraph('The CLOCS CMSC adaptation averages 0.2877 F1 and 0.8264 AUROC. SimCLR domain minus CLOCS has AUROC difference −0.0119, with adjusted p = 0.0137. Its F1 difference is −0.0272 with adjusted p = 0.0684; the unadjusted bootstrap interval excludes zero, but the adjusted exact test does not meet 0.05. We report both procedures rather than equating an unadjusted interval with a multiplicity-adjusted decision.')
figure('performance.png', 'Figure 1 Frozen and fully fine-tuned shared-backbone performance across ten seeds with pointwise seed confidence intervals')

heading('4 3 Per-class performance',2)
table('Table 6 SimCLR domain class performance at threshold 0.5 across ten seeds',
      ['Class','Positive support','Recall mean ± SD','Precision mean ± SD','F1 mean ± SD','Specificity mean ± SD'],[
      [c,DATA['methods']['simclr_domain']['support'][c],*[estimate(DATA['methods']['simclr_domain']['per_class'][c][m]) for m in ['sensitivity','precision','f1','specificity']]] for c in ['NORM','CD','HYP','MI','STTC']], [0.6,.85,1.33,1.33,1.33,1.33])
paragraph('The held-out test set contains 2,158 eligible records. HYP has the smallest positive support (262), not CD (496). NORM has 963, MI 550 and STTC 521 positive records. Supports sum to 2,792 because diagnoses co-occur. Table 6 averages class metrics over seeds; macro-F1 is the unweighted mean across five classes, and a support-weighted average is not substituted for it. Detailed class AUROC/PR-AUC and seed-level results accompany the analysis CSV files.')

heading('4 4 Leave-one-out ablation',2)
table('Table 7 SimCLR omissions and paired full-policy-minus-omission effects',
      ['Omitted transformation','F1 mean','ΔF1 95% CI','Holm p F1','ΔAUROC','Holm p AUC'],[
      [name.replace('_',' '),f"{S['summaries']['simclr_without_'+name]['f1_macro']['mean']:.4f}",
       f"{S['paired_comparisons']['simclr_domain__minus__simclr_without_'+name]['f1_macro']['mean']:+.4f} "+interval(S['paired_comparisons']['simclr_domain__minus__simclr_without_'+name]['f1_macro']),
       f"{S['paired_comparisons']['simclr_domain__minus__simclr_without_'+name]['f1_macro']['p_holm']:.4f}",
       f"{S['paired_comparisons']['simclr_domain__minus__simclr_without_'+name]['auroc_macro']['mean']:+.4f}",
       f"{S['paired_comparisons']['simclr_domain__minus__simclr_without_'+name]['auroc_macro']['p_holm']:.4f}"]
      for name in ['frequency_warp','mixup','bandpass','cutmix','motion','channel_noise','temporal_dropout']], [1.35,.65,2.1,.8,.9,.97])
paragraph('Mixup increases full-policy F1 by 0.0323 relative to its omission (adjusted p = 0.0137), without a detectable AUROC benefit. Conversely, the full policy is 0.0277 F1 below the no-CutMix condition (adjusted p = 0.0137) and 0.0041 AUROC below it (adjusted p = 0.0410). Bandpass omission has higher mean F1, but its adjusted p = 0.3223 does not establish an improvement. The other omissions are not distinguishable after adjustment; absence of significance is not proof that they are redundant. These are conditional effects within this library, not estimates of standalone transformation efficacy or pairwise interactions. A reduced policy suggested by these findings would require separate validation rather than selecting a new headline model on these test outcomes.')

heading('4 5 Thresholds calibration and frozen BYOL regression',2)
table('Table 8 Threshold and calibration diagnostics across ten seeds',
      ['Method','Fixed F1','Validation-tuned F1 mean ± SD','Brier mean','ECE mean'],[
      [NAMES[n],f"{S['summaries'][n]['f1_macro']['mean']:.4f}",estimate(DATA['methods'][n]['tuned']['f1_macro']),f"{S['summaries'][n]['brier_macro']['mean']:.4f}",f"{S['summaries'][n]['ece_macro']['mean']:.4f}"] for n in DISPLAY_METHODS], [1.77,.85,2.05,1.05,1.05])
paragraph('Validation-selected thresholds increase mean F1 to 0.5880 for SimCLR domain and 0.6237 for BYOL domain, compared with 0.6358 supervised. This descriptive secondary analysis shows that the operating threshold explains part of the large fixed-threshold F1 deficit. It does not erase the lower AUROC or demonstrate better calibration. BYOL domain has mean Brier score 0.1447 and ECE 0.1457, compared with supervised 0.1231 and 0.0773; lower values are preferable for both metrics.')
paragraph('BYOL domain minus supervised AUROC is −0.0130, with paired 95% CI [−0.0170, −0.0091] and adjusted p = 0.0117. Thus the regression is distinguishable from training-seed variation under the tested protocol. Its validation representation has mean effective rank 104.6 of 256 dimensions and nonzero feature variance, which does not suggest complete constant-output collapse. Effective rank is a diagnostic, not proof of discriminative quality. Neither the AUROC regression nor these diagnostics identifies the absence of negative pairs as a causal mechanism; BYOL actually exceeds SimCLR here.')
paragraph('Full domain BYOL fine-tuning has mean AUROC 0.8594, Brier score 0.1173 and ECE 0.0617. The Brier/ECE means improve descriptively over frozen BYOL and supervised learning. Calibration contrasts are outside the original statistical plan; separately reported post hoc inference distinguishes the evidence for each comparison. The frozen AUROC deficit therefore does not persist as a demonstrated fine-tuning deficit, but the fine-tuned AUROC gain also does not survive multiplicity adjustment. Operating-point F1 and calibration must be assessed separately.')

heading('4 6 Patient sampling uncertainty',2)
table('Table 9 Patient cluster and joint patient and seed confidence intervals',
      ['Comparison','Macro F1 joint 95% CI','Macro AUROC joint 95% CI'], [
      [label, '['+', '.join(f'{v:+.4f}' for v in PATIENT['paired_intervals'][pair]['f1_macro']['patients_and_seeds_ci95'])+']',
       '['+', '.join(f'{v:+.4f}' for v in PATIENT['paired_intervals'][pair]['auroc_macro']['patients_and_seeds_ci95'])+']']
      for pair,label in [('simclr_domain__minus__simclr_generic','SimCLR domain minus generic'),
                         ('byol_domain__minus__byol_generic','BYOL domain minus generic'),
                         ('simclr_domain__minus__supervised','SimCLR domain minus supervised'),
                         ('byol_domain__minus__supervised','BYOL domain minus supervised')]], [2.77,2,2])
paragraph('Patient uncertainty is estimated using 2,000 paired cluster-bootstrap draws from the 1,877 test subjects, retaining all 2,158 records and repeated measurements within each sampled subject. Joint intervals additionally resample the ten independently trained seeds and describe the mean single-model metric, rather than an ensemble. All methods share identical patient and seed draws. These intervals remain conditional on training and validation subjects and fixed thresholds. BYOL domain minus supervised AUROC remains negative under joint patient and seed uncertainty. Intervals are pointwise and do not quantify population shift or prospective clinical performance.')

heading('4 7 Training budget and convergence',2)
table('Table 10 Measured mean training duration on Apple M5 Pro',
      ['Method','Pretraining epochs','Pretraining seconds mean','Downstream seconds mean'],[
      [NAMES[n],'None' if n in ['supervised','clocs_reference_supervised'] else ('20 reused' if 'finetuned' in n else '20'),
       'Reused encoder' if 'finetuned' in n else f"{DATA['methods'][n]['times']['pretrain_seconds']['mean']:.1f}",
       f"{DATA['methods'][n]['times']['downstream_seconds']['mean']:.1f}"] for n in DISPLAY_METHODS], [2.5,1.15,1.65,1.47])
paragraph('Histories retain all pretraining losses and validation metrics, and convergence plots are included with the saved analysis. Twenty epochs are a matched prespecified budget, not a claim of convergence or optimality. Domain policies cost more wall time than generic policies, partly because transformations run on the CPU before MPS transfer. CLOCS uses shorter windows. Timings describe this machine and protocol; they do not imply a hardware-independent computational ranking. Best epochs vary by seed and are selected by validation F1; no universal epoch-four stopping rule is justified.')

if ADDITIONAL is not None:
    heading('4 8 BYOL augmentation ablations',2)
    table('Table 11 BYOL domain augmentation omissions across ten seeds',
      ['Omitted transformation','Macro F1 mean ± SD','ΔF1','Holm p F1','ΔAUROC','Holm p AUC'],[
          [name.replace('_',' '),estimate(ADDITIONAL['summaries']['byol_without_'+name]['f1_macro']),
           f"{ADDITIONAL['paired_comparisons']['byol_domain__minus__byol_without_'+name]['f1_macro']['mean']:+.4f}",
           f"{ADDITIONAL['paired_comparisons']['byol_domain__minus__byol_without_'+name]['f1_macro']['p_holm']:.4f}",
           f"{ADDITIONAL['paired_comparisons']['byol_domain__minus__byol_without_'+name]['auroc_macro']['mean']:+.4f}",
           f"{ADDITIONAL['paired_comparisons']['byol_domain__minus__byol_without_'+name]['auroc_macro']['p_holm']:.4f}"]
          for name in ['frequency_warp','mixup','bandpass','cutmix','motion','channel_noise','temporal_dropout']], [1.5,1.67,.9,.9,.9,.9])
    paragraph('BYOL omission effects are estimated using the same seven transformations, ten seeds and labelled subsets as the SimCLR analysis. Positive full-minus-omitted differences favour retaining the transformation. Holm adjustment treats the seven BYOL comparisons as a separate prespecified family. These effects are conditional on the remaining policy and do not identify interactions.')
    paragraph('For BYOL, bandpass variation improves F1 by 0.0315 and AUROC by 0.0047 (both adjusted p = 0.0137). Frequency warping improves F1 by 0.0228 (adjusted p = 0.0352), without a detectable AUROC effect. Retaining mixup reduces AUROC by 0.0043, and temporal dropout reduces it by 0.0023 (both adjusted p = 0.0137), without detectable F1 effects. The other BYOL differences do not survive adjustment. Thus the SimCLR mixup benefit cannot be generalized to BYOL, and augmentation choices involve metric-specific trade-offs.')
    figure('augmentation_effects.png', 'Figure 2 Full-policy-minus-omission effects with pointwise seed confidence intervals and separate Holm decisions for each framework and metric')
    heading('4 9 Full encoder fine tuning',2)
    fine_names=[n+'_finetuned' for n in METHODS[1:]]
    table('Table 12 Full SSL fine tuning across ten seeds',
          ['Method','Macro F1 mean ± SD','Macro AUROC mean ± SD','F1 minus supervised'],[
          [NAMES[name.removesuffix('_finetuned')],estimate(ADDITIONAL['summaries'][name]['f1_macro']),
           estimate(ADDITIONAL['summaries'][name]['auroc_macro']),
           f"{ADDITIONAL['paired_comparisons'][name+'__minus__supervised']['f1_macro']['mean']:+.4f}"]
          for name in fine_names], [1.8,1.8,1.8,1.37])
    paragraph('Full fine-tuning reuses each seed-specific pretrained encoder, verified by checksum and parent fingerprint. The labelled, validation and test record IDs match its frozen counterpart. Every encoder parameter is trainable; Adam at 0.001, focal loss, oversampling, 20 downstream epochs and validation macro-F1 checkpoint selection match supervised training. This comparison isolates downstream encoder training within the evaluated pretraining and learning schedule. Seed-paired tests for fine-tuned-minus-frozen and fine-tuned-minus-supervised comparisons form a separate prespecified family. Hyperparameters are not selected using test results.')
    paragraph('All five fine-tuned conditions improve F1 and AUROC over their respective frozen counterparts (all adjusted p = 0.0195). Domain SimCLR and BYOL exceed supervised F1 by 0.0259 and 0.0265, with joint patient-and-seed intervals excluding zero. Their AUROC gains do not survive Holm adjustment, and the joint AUROC intervals include zero. Shared-backbone CLOCS fine-tuning instead improves AUROC by 0.0096 (adjusted p = 0.0195), without a detectable F1 advantage. Generic BYOL has a positive pointwise joint F1 interval but adjusted p = 0.0645; the pointwise interval is not a multiplicity-adjusted decision. Fine-tuned domain-minus-generic mean F1 differences are 0.0193 for SimCLR and 0.0127 for BYOL, reported descriptively outside the prespecified paired test families.')
    for name in fine_names:
        statistics=ADDITIONAL['paired_comparisons'][name+'__minus__supervised']
        pieces=[]
        for metric,label in [('f1_macro','macro-F1'),('auroc_macro','macro-AUROC')]:
            values=statistics[metric]
            joint=PATIENT['paired_intervals'][name+'__minus__supervised'][metric]['patients_and_seeds_ci95']
            pieces.append(f"{label} difference {values['mean']:+.4f}, seed 95% CI {interval(values)}, adjusted p {values['p_holm']:.4f}, joint patient and seed 95% CI [{joint[0]:+.4f}, {joint[1]:+.4f}]")
        paragraph(NAMES[name.removesuffix('_finetuned')]+' full fine-tuning minus supervised: '+ '; '.join(pieces)+'. These signed differences report gains and deficits under the declared learning schedule.')
    figure('validation_convergence.png', 'Figure 3 Validation convergence at the matched downstream epoch budget with sample standard deviation across ten seeds')
    heading('4 10 CLOCS reference architecture',2)
    reference_names=['clocs_reference_supervised','clocs_reference_cmsc','clocs_reference_finetuned']
    table('Table 13 CLOCS reference architecture on the common diagnostic task',
          ['Condition','Macro F1 mean ± SD','Macro AUROC mean ± SD'],[
          [{'clocs_reference_supervised':'Supervised', 'clocs_reference_cmsc':'Frozen CMSC', 'clocs_reference_finetuned':'Fully fine-tuned CMSC'}[name],estimate(ADDITIONAL['summaries'][name]['f1_macro']),
           estimate(ADDITIONAL['summaries'][name]['auroc_macro'])] for name in reference_names], [2.27,2.25,2.25])
    paragraph('The reference encoder follows the author implementation: three kernel-seven, stride-three convolution blocks with channels 1–4–16–32, batch normalization, ReLU, max pooling and dropout 0.1, followed by a 320-to-256 linear embedding. CMSC pairs adjacent five-second lead-II segments at native 500 Hz with the patient-specific objective and no additional augmentation. The downstream input is the first five seconds. Frozen and trainable pretrained encoders are compared with a random-initialized control using identical architecture, input, labels and downstream settings. This evaluates the reference architecture on the present five-superclass task; it does not reproduce the original paper’s different downstream datasets or label definitions.')
    paragraph('Frozen reference CMSC improves F1 over its architecture-matched supervised control by 0.0784, but reduces AUROC by 0.0733. Full reference fine-tuning improves F1 by 0.0252 and AUROC by 0.0130 over that control; against its frozen counterpart it reduces F1 by 0.0531 and improves AUROC by 0.0863. All six metric tests have adjusted p = 0.0059, and their joint patient-and-seed intervals exclude zero. These opposing changes show why threshold F1 alone cannot establish global discrimination or calibration. Cross-backbone comparisons also change lead count, sampling rate and window length.')
    for pair in ['clocs_reference_cmsc__minus__clocs_reference_supervised', 'clocs_reference_finetuned__minus__clocs_reference_supervised', 'clocs_reference_finetuned__minus__clocs_reference_cmsc']:
        pieces=[]
        for metric,label in [('f1_macro','F1'),('auroc_macro','AUROC')]:
            values=ADDITIONAL['paired_comparisons'][pair][metric]
            joint=PATIENT['paired_intervals'][pair][metric]['patients_and_seeds_ci95']
            pieces.append(f"{label} {values['mean']:+.4f}, seed 95% CI {interval(values)}, joint 95% CI [{joint[0]:+.4f}, {joint[1]:+.4f}]")
        paragraph(pair.replace('clocs_reference_','').replace('__minus__',' minus ').replace('_',' ')+': '+ '; '.join(pieces)+'.')

POSTHOC = ROOT/'results/extensions/analysis/paired_inference.json'
if POSTHOC.exists():
    ph=json.loads(POSTHOC.read_text())
    heading('4 11 Paired policy and calibration inference',2)
    paragraph('The following analyses of existing predictions are post hoc and do not change the original inferential families. Holm correction is applied jointly to four fine-tuned domain-minus-generic F1 and AUROC contrasts, and jointly to ten calibration contrasts. Joint intervals resample test subjects and training seeds and remain pointwise.')
    table('Table 14 Post hoc full fine tuning domain minus generic policy comparisons',
          ['Framework and metric','Mean difference','Joint 95% CI','Adjusted p'],[
          [r['method_a'].split('_')[0].upper()+' '+('F1' if r['metric']=='f1_macro' else 'AUROC'),
           f"{r['mean_difference']:+.4f}",f"[{r['joint_ci95'][0]:+.4f}, {r['joint_ci95'][1]:+.4f}]",f"{r['p_holm']:.4f}"]
          for r in ph['comparisons'] if r['family']=='downstream_policy'],[2.0,1.1,2.5,1.17])
    paragraph('None of the four fine-tuned policy contrasts meets 0.05 after joint Holm adjustment. Frozen BYOL domain has worse Brier score and ECE than supervised learning (both adjusted p = 0.0195). Full fine-tuning improves both relative to frozen BYOL (both adjusted p = 0.0195), but its calibration advantages over supervised learning are not established after adjustment. Lower Brier score and ECE indicate better calibration scores; these differences do not isolate an objective-level causal mechanism.')

EXTENSION_COMPLETE=ROOT/'results/extensions/analysis/completion.json'
if EXTENSION_COMPLETE.exists():
    panels={name:json.loads((ROOT/'results/extensions/analysis'/name/'summary.json').read_text()) for name in ('matched_generic','label_efficiency','external_transfer')}
    heading('4 12 Gate and RMS matched generic controls',2)
    report=panels['matched_generic']
    table('Table 15 Matched generic controls across ten training seeds', ['Condition','F1 mean ± SD','AUROC mean ± SD'],[
        [NAMES.get(name,name.replace('_',' ').replace('simclr','SimCLR').replace('byol','BYOL')),estimate(v['f1_macro']),estimate(v['auroc_macro'])] for name,v in report['summaries'].items()],[2.77,2.,2.])
    paragraph('Matching increases frozen generic F1 by 0.1339 for SimCLR and 0.1422 for BYOL versus the simpler generic policy (both adjusted p = 0.0156). SimCLR domain minus matched generic is −0.0250 F1 (adjusted p = 0.0586) and −0.0120 AUROC (adjusted p = 0.0156). BYOL retains domain advantages of 0.0471 F1 and 0.0087 AUROC (adjusted p = 0.0156 and 0.0293). Full fine-tuning does not establish domain advantages over matched generic controls in either framework. The contrast is sensitive to the generic control construction; this experiment does not isolate gate count from RMS strength or operator choice.')
    uncertainty=json.loads((ROOT/'results/extensions/analysis/matched_generic/analysis/patient_uncertainty.json').read_text())
    rows=[]
    for framework in ('simclr','byol'):
        pair=framework+'_domain__minus__'+framework+'_generic_matched'
        for metric in ('f1_macro','auroc_macro'):
            v=report['paired_comparisons'][pair][metric];ci=uncertainty['paired_intervals'][pair][metric]['patients_and_seeds_ci95']
            rows.append([framework.upper()+' '+('F1' if metric=='f1_macro' else 'AUROC'),f"{v['mean']:+.4f}",f"[{ci[0]:+.4f}, {ci[1]:+.4f}]",f"{v['p_holm']:.4f}"])
    table('Table 16 Frozen domain minus matched generic contrasts', ['Framework and metric','Difference','Joint 95% CI','Adjusted p'],rows,[2.,1.1,2.5,1.17])
    paragraph('Figure 4 Matched augmentation controls. Error bars are pointwise 95% bootstrap intervals for the seed mean; inferential decisions use the paired panel tests.','Caption').paragraph_format.keep_with_next=True
    doc.add_picture(str(ROOT/'results/extensions/analysis/matched_generic.png'),width=Inches(6.7))
    heading('4 13 Label efficiency',2)
    report=panels['label_efficiency']
    rows=[]
    for parent in ('supervised','simclr_generic','simclr_domain','byol_generic','byol_domain'):
        values=[]
        for fraction in (1,5,10,20,50):
            name=f'{parent}_labels_{fraction:02d}' if fraction!=10 else parent+('_finetuned' if parent!='supervised' else '')
            v=report['summaries'][name];values.append(f"{v['f1_macro']['mean']:.4f} / {v['auroc_macro']['mean']:.4f}")
        rows.append([NAMES[parent]]+values)
    table('Table 17 Label efficiency: mean F1 / AUROC across ten seeds', ['Condition','1%','5%','10%','20%','50%'],rows,[1.77,1.,1.,1.,1.,1.])
    paragraph('Supervised mean F1 rises from 0.4804 at 1% to 0.6409 at 50%. Domain SSL often has higher mean F1 than its generic counterpart, but no F1 or AUROC contrast survives the 30-comparison panel adjustment (minimum adjusted p = 0.0586). At 10%, SimCLR and BYOL domain minus supervised F1 have adjusted p = 0.1055 and 0.0586 in this broader family, versus 0.0195 in the dedicated fine-tuning family in Section 4 9. The estimates are identical; different declared hypothesis families explain the different adjusted decisions. These results do not establish a consistent label-efficiency advantage across fractions. Complete seed statistics and pointwise patient/joint intervals accompany the paired analysis tables.')
    paragraph('Figure 5 Full encoder training across labelled fractions. Lines show means and shaded bands pointwise 95% seed-mean intervals. The horizontal axis is logarithmic.','Caption').paragraph_format.keep_with_next=True
    doc.add_picture(str(ROOT/'results/extensions/analysis/label_efficiency.png'),width=Inches(6.7))
    heading('4 14 MIT-BIH annotation-window transfer',2)
    report=panels['external_transfer']
    table('Table 18 MIT-BIH transfer across ten training seeds', ['Initialization','F1 mean ± SD','AUROC mean ± SD'],[
        [('Supervised BCE' if name=='mitbih_supervised' else NAMES.get(name.removeprefix('mitbih_'),name)),estimate(v['f1_macro']),estimate(v['auroc_macro'])] for name,v in report['summaries'].items()],[2.77,2.,2.])
    paragraph('Supervised mean F1/AUROC is 0.7480/0.8531. Generic SimCLR has the highest means, 0.7647/0.8668, while domain SimCLR is lower, 0.7320/0.8478. BYOL generic and domain reach 0.7520/0.8572 and 0.7561/0.8610. No contrast survives the six-comparison panel adjustment: all F1 adjusted p-values are 1.0000; AUROC adjusted p-values range from 0.7383 to 1.0000. Generic SimCLR minus supervised joint intervals include zero for F1 [−0.0454, 0.0959] and AUROC [−0.0327, 0.0787]. With nine held-out subjects, patient uncertainty is substantial. The results provide no detected transfer advantage and do not establish equivalence.')
    paragraph('Figure 6 Binary annotation-window transfer. Error bars are pointwise seed-mean intervals on the single fixed MIT-BIH subject partition.','Caption').paragraph_format.keep_with_next=True
    doc.add_picture(str(ROOT/'results/extensions/analysis/external_transfer.png'),width=Inches(6.7))

heading('5 Discussion')
if ADDITIONAL is not None:
    paragraph('All 570 runs are complete. The gate/RMS matching, five-fraction label-efficiency and MIT-BIH transfer panels complement the encoder-adaptation, omission and CLOCS controls, with paired seed, patient and joint uncertainty. No model settings were selected from these test outcomes.')

paragraph('The simpler generic controls show domain-policy improvements within frozen SimCLR and BYOL, but the gate/RMS matched controls materially qualify attribution. Frozen SimCLR domain AUROC is below matched generic AUROC, while frozen BYOL retains domain advantages. Neither framework establishes a fine-tuned advantage over matched generic controls. The omission results further show that individual effects depend on the full library and metric: mixup benefits SimCLR fixed-threshold F1, while CutMix reduces it. Physiological motivation alone is insufficient evidence for choosing an augmentation.')
paragraph('The supervised comparator is strongest in the frozen-encoder comparison. Its trainable encoder can adapt features to the labelled task, whereas SSL linear evaluation can adapt only the head. The secondary threshold analysis narrows F1 differences but leaves ranking and calibration differences. End-to-end SSL fine-tuning is evaluated separately with the same labelled subsets, downstream optimizer, epoch budget and validation checkpoint criterion. Domain SimCLR and BYOL gain fixed-threshold F1 after full fine-tuning, while their AUROC gains against supervised learning are not detectable after adjustment. Fine-tuning improves all five shared-backbone SSL conditions over their frozen counterparts on both metrics (all adjusted p = 0.0195). These findings show that downstream encoder adaptation changes the comparative conclusion; superiority is metric- and protocol-dependent.')
paragraph('Frozen BYOL’s significant AUROC regression deserves equal prominence to the within-SSL augmentation gains. These observations do not establish that negative pairs make SimCLR globally more discriminative. No objective-level intervention isolates that mechanism, and the present BYOL results exceed SimCLR. We report the observed ranking, calibration and representation diagnostics while leaving the cause unresolved.')

heading('5 1 Limitations',2)
paragraph('The diagnostic experiment uses one main PTB-XL cohort, five labelled fractions and one learning schedule. Gate/RMS matching controls two augmentation quantities but does not match spectral, morphological or clinical fidelity. The matched policy still represents only one generic construction. MIT-BIH provides corroborative binary-window transfer on a previously evaluated fixed subject partition, not untouched external confirmation, an AAMI beat benchmark or ten independent patient partitions. Neither the label-efficiency panel nor transfer panel establishes a general SSL advantage. Device invariance, comorbidity simulation and clinical fidelity of mixed ECGs remain unvalidated. Complete seven-omission studies in both frameworks estimate effects conditional on the remaining library and cannot identify interactions. Follow-up families were declared after original outcomes were known; post hoc policy/calibration contrasts remain labelled as such. Bootstrap intervals quantify seed and held-out-patient uncertainty conditional on observed training/validation subjects; they are pointwise, not simultaneous, and do not quantify population shift or prospective clinical performance. Equal epochs do not equalize FLOPs, view durations or elapsed time; measured runtime profiles are not time-matched performance evidence. Shared-backbone and native-reference CLOCS comparisons have different inputs and architectures; matched supervised controls support within-architecture inference, while cross-architecture rankings retain these confounds. Report-based SSL requires inputs outside this waveform-only study.')

heading('6 Conclusion')
paragraph('Across 570 runs, augmentation effects depend on objective, control strength, metric and downstream adaptation. Domain policies improve frozen representations over a simpler generic policy, but matched gate/RMS controls remove a universal domain advantage and reverse the frozen SimCLR AUROC comparison. Full domain fine-tuning improves 10% macro-F1 over supervised learning in the dedicated family, without an adjusted AUROC advantage; broader label-efficiency and MIT-BIH families detect no advantage. Conditional omission effects include both benefits and harms. These results support protocol-specific conclusions and further clinical validation rather than general SSL superiority.')

heading('Ethical statement and availability')
paragraph('This study uses public, de-identified PTB-XL and MIT-BIH datasets [7, 9–11] and collects no new patient information. Models remain research prototypes without prospective diagnostic or screening validation. Dataset use is subject to the terms accompanying the cited releases. The project repository is https://github.com/Tumo505/SSL-for-ECG-classification. Local outputs include provenance, checkpoints, predictions and analysis tables; public release of outputs and weights is not yet established.')

heading('References')
refs=[
 '[1] Chen T, Kornblith S, Norouzi M, Hinton G. A Simple Framework for Contrastive Learning of Visual Representations. ICML, PMLR 119:1597–1607, 2020. https://proceedings.mlr.press/v119/chen20j.html',
 '[2] Grill JB et al. Bootstrap your own latent: A new approach to self-supervised learning. NeurIPS, 2020. https://arxiv.org/abs/2006.07733',
 '[3] Kiyasseh D, Zhu T, Clifton DA. CLOCS: Contrastive Learning of Cardiac Signals Across Space, Time, and Patients. ICML, PMLR 139:5606–5615, 2021. https://proceedings.mlr.press/v139/kiyasseh21a.html. Author code: https://github.com/danikiyasseh/CLOCS',
 '[4] Li H et al. Fine-grained Contrastive Learning for ECG-Report Alignment with Waveform Enhancement. 2025. https://arxiv.org/abs/2505.11939',
 '[5] Srivastava A et al. Channel self-attention deep learning framework for multi-cardiac abnormality diagnosis from varied-lead ECG signals. Computing in Cardiology, 2021. doi:10.23919/CinC53138.2021.9662886',
 '[6] Wagner P et al. PTB-XL, a large publicly available electrocardiography dataset. Scientific Data 7:154, 2020. doi:10.1038/s41597-020-0495-6',
 '[7] Wagner P et al. PTB-XL, a large publicly available electrocardiography dataset, version 1.0.3. PhysioNet, 2022. https://physionet.org/content/ptb-xl/1.0.3/',
 '[8] Lin TY et al. Focal Loss for Dense Object Detection. ICCV, 2017. https://openaccess.thecvf.com/content_iccv_2017/html/Lin_Focal_Loss_for_ICCV_2017_paper.html']
refs.extend(['[9] Moody G, Mark R. MIT-BIH Arrhythmia Database, version 1.0.0. PhysioNet, 2005. doi:10.13026/C2F305. https://physionet.org/content/mitdb/1.0.0/', '[10] Moody GB, Mark RG. The impact of the MIT-BIH Arrhythmia Database. IEEE Engineering in Medicine and Biology Magazine 20(3):45–50, 2001. doi:10.1109/51.932724', '[11] Pollard T et al. PhysioNet as a global platform for biomedical research. Nature Health 1:792–795, 2026. doi:10.1038/s44360-026-00096-z'])
for ref in refs:
    p=paragraph(ref)
    p.paragraph_format.space_after=Pt(3)
    for run in p.runs:run.font.size=Pt(9.5)
footer=section.footer.paragraphs[0]
footer.alignment=WD_ALIGN_PARAGRAPH.CENTER
field=OxmlElement('w:fldSimple');field.set(qn('w:instr'),'PAGE');footer._p.append(field)
doc.core_properties.title='Evaluating domain adaptive augmentations for self supervised ECG classification'
doc.core_properties.subject='PTB-XL and MIT-BIH experiments with complete paired multi-seed results'
for element in list(doc.core_properties._element):
    if element.tag == qn('cp:revision'):
        doc.core_properties._element.remove(element)
doc.save(ROOT/'artifacts/Deep Learning Indaba.docx')
print(ROOT/'artifacts/Deep Learning Indaba.docx')
